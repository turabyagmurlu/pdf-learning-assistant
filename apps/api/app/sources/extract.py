"""PDF disindaki kaynaklari "bolum"lere cevirir.

Her tur ayni ciktiyi uretir: [{page_number, title, text, table?}]
- page_number: atif birimi (Word/Markdown'da bolum, sunumda slayt, tabloda blok)
- text: dizinlenecek ve aranacak duz metin
- table: (varsa) okuyucuda tablo olarak gostermek icin {header, rows}
Hepsi yerel calisir; yapay zeka kotasi harcamaz.
"""
import csv
import io
import re
import zipfile
from app.core.errors import AppError

WORDS_PER_SECTION = 600
ROWS_PER_BLOCK = 40

# uzanti -> (tur, okuyucuda gorunen birim)
KINDS = {
    "docx": ("docx", "böl."), "md": ("md", "böl."), "markdown": ("md", "böl."),
    "txt": ("txt", "böl."), "rtf": ("rtf", "böl."), "xlsx": ("xlsx", "tablo"),
    "xlsm": ("xlsx", "tablo"), "csv": ("csv", "tablo"), "tsv": ("csv", "tablo"),
    "pptx": ("pptx", "slayt"), "epub": ("epub", "böl."), "html": ("html", "böl."),
    "htm": ("html", "böl."),
}
UNIT = {"docx": "böl.", "md": "böl.", "txt": "böl.", "rtf": "böl.", "xlsx": "tablo", "csv": "tablo",
        "pptx": "slayt", "epub": "böl.", "html": "böl.", "web": "böl.", "text": "böl.", "audio": "böl."}
UNSUPPORTED_OLD = {"doc": "Word (.doc) eski biçim; Word'de 'Farklı kaydet → .docx' yapıp yükle.",
                   "xls": "Excel (.xls) eski biçim; Excel'de 'Farklı kaydet → .xlsx' yapıp yükle.",
                   "ppt": "PowerPoint (.ppt) eski biçim; '.pptx' olarak kaydedip yükle."}


def kind_of(filename: str) -> str | None:
    ext = (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""
    if ext in UNSUPPORTED_OLD:
        raise AppError(UNSUPPORTED_OLD[ext])
    k = KINDS.get(ext)
    return k[0] if k else None


# ------------------------------------------------------------------ yardimci
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1254", "iso-8859-9", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "ignore")


def _clean(t: str) -> str:
    t = (t or "").replace("\r", "").replace("\xa0", " ")
    t = re.sub(r"[ \t]+", " ", t)
    return re.sub(r"\n{3,}", "\n\n", t).strip()


def _split_long(title: str | None, text: str, limit: int = WORDS_PER_SECTION) -> list[tuple[str | None, str]]:
    """Uzun bir bolumu paragraf sinirlarinda ~limit kelimelik parcalara boler."""
    paras = []
    for p in re.split(r"\n\s*\n|\n", text):
        if not p.strip():
            continue
        words = p.split()
        if len(words) > limit:                    # tek dev paragraf: kelime sinirinda parcala
            for i in range(0, len(words), limit):
                paras.append(" ".join(words[i:i + limit]))
        else:
            paras.append(p)
    out, cur, n, part = [], [], 0, 1
    for p in paras:
        w = len(p.split())
        if cur and n + w > limit:
            out.append((title if part == 1 else f"{title or 'Devam'} ({part})", "\n".join(cur)))
            cur, n, part = [], 0, part + 1
        cur.append(p); n += w
    if cur:
        out.append((title if part == 1 else f"{title or 'Devam'} ({part})", "\n".join(cur)))
    return out


def _to_pages(sections: list[tuple[str | None, str]]) -> list[dict]:
    pages = []
    for title, text in sections:
        text = _clean(text)
        if not text:
            continue
        for t, chunk in _split_long(title, text):
            body = f"{t}\n{chunk}" if t and not chunk.startswith(t) else chunk
            pages.append({"page_number": len(pages) + 1, "title": t, "text": body})
    return pages


def _by_headings(lines: list[tuple[int, str]]) -> list[tuple[str | None, str]]:
    """(seviye, satir) listesini baslik bazli bolumlere cevirir. seviye 0 = normal paragraf."""
    secs: list[tuple[str | None, list[str]]] = []
    cur_title, cur = None, []
    for lvl, line in lines:
        if lvl and lvl <= 3 and line.strip():
            if cur or cur_title:
                secs.append((cur_title, cur))
            cur_title, cur = line.strip(), []
        else:
            cur.append(line)
    if cur or cur_title:
        secs.append((cur_title, cur))
    return [(t, "\n".join(c)) for t, c in secs]


# ------------------------------------------------------------------ turler
def _docx(data: bytes):
    import docx  # python-docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(io.BytesIO(data))
    lines: list[tuple[int, str]] = []
    body = d.element.body
    for el in body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, d)
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"(Heading|Başlık|Title|Konu Başlığı)\s*(\d*)", style, re.I)
            lvl = int(m.group(2) or 1) if m else 0
            if p.text.strip():
                lines.append((lvl, p.text))
        elif tag == "tbl":
            t = Table(el, d)
            for row in t.rows:
                cells = []
                for c in row.cells:
                    v = c.text.strip()
                    if not cells or cells[-1] != v:          # birlesik hucre tekrarini at
                        cells.append(v)
                if any(cells):
                    lines.append((0, " | ".join(cells)))
    title = (d.core_properties.title or "").strip() or None
    return _to_pages(_by_headings(lines)), title


def _markdown(data: bytes):
    text = _decode(data)
    lines = []
    in_code = False
    for raw in text.splitlines():
        if raw.strip().startswith("```"):
            in_code = not in_code
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", raw) if not in_code else None
        lines.append((len(m.group(1)), m.group(2)) if m else (0, raw))
    first = next((l for lv, l in lines if lv == 1), None)
    return _to_pages(_by_headings(lines)), first


def _txt(data: bytes):
    text = _clean(_decode(data))
    first = next((l.strip() for l in text.splitlines() if l.strip()), "")[:120] or None
    return _to_pages([(None, text)]), first


def _rtf(data: bytes):
    raw = _decode(data)
    try:
        from striprtf.striprtf import rtf_to_text
        text = rtf_to_text(raw)
    except Exception:  # noqa - kaba temizlik
        text = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", raw)
    return _to_pages([(None, text)]), None


def _table_pages(sheet: str, rows: list[list[str]], start_no: int) -> list[dict]:
    rows = [[("" if c is None else str(c)).strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    # bos sutunlari at
    keep = [i for i in range(width) if any(r[i] for r in rows)]
    rows = [[r[i] for i in keep] for r in rows]
    header = rows[0]
    has_header = sum(1 for h in header if h and not re.fullmatch(r"[\d.,\-]+", h)) >= max(1, len(header) // 2)
    if not has_header:
        header = [f"Sütun {i + 1}" for i in range(len(rows[0]))]
        data = rows
    else:
        data = rows[1:] or []
        header = [h or f"Sütun {i + 1}" for i, h in enumerate(header)]
    pages = []
    stats = _table_stats(sheet, header, data, 2 if has_header else 1)
    if stats:
        pages.append({"page_number": start_no, "title": f"{sheet} · özet istatistik", "text": stats})
    for b in range(0, max(1, len(data)), ROWS_PER_BLOCK):
        block = data[b:b + ROWS_PER_BLOCK]
        first_row = b + (2 if has_header else 1)
        last_row = first_row + len(block) - 1
        title = f"{sheet} · satır {first_row}–{last_row}" if block else sheet
        # anlam aramasi icin her satir "Baslik: deger" bicimde
        lines = [f"Tablo: {sheet}. Sütunlar: {', '.join(header)}."]
        for i, r in enumerate(block):
            kv = "; ".join(f"{h}: {v}" for h, v in zip(header, r) if v)
            if kv:
                lines.append(f"Satır {first_row + i} — {kv}")
        pages.append({"page_number": start_no + len(pages), "title": title, "text": "\n".join(lines),
                      "table": {"header": header, "rows": block, "first_row": first_row}})
    return pages


def _num(v: str):
    s = (v or "").strip().replace("−", "-").replace("%", "")
    if not s:
        return None
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s):      # 1.234.567,89 (TR)
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d+,\d+", s):                       # 12,5
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _fmt(x: float) -> str:
    return f"{x:,.2f}".rstrip("0").rstrip(".").replace(",", "X").replace(".", ",").replace("X", ".")


def _table_stats(sheet: str, header: list[str], data: list[list[str]], first_row: int) -> str | None:
    """Tablonun yerel ozeti: 'ortalama/en yuksek/kac tane' sorularini dogru cevaplamak icin.
    Yapay zeka satir satir toplama yapamaz; bu ozet hesap hatasini onler."""
    if len(data) < 3:
        return None
    lines = [f"Tablo: {sheet} — özet istatistik ({len(data)} satır, {len(header)} sütun). "
             f"Sütunlar: {', '.join(header)}."]
    label_col = 0
    for ci, h in enumerate(header):
        col = [r[ci] for r in data if ci < len(r)]
        filled = [v for v in col if v]
        nums = [(_num(v), ri) for ri, v in enumerate(col) if _num(v) is not None]
        if filled and len(nums) >= 0.8 * len(filled) and len(nums) >= 2:
            vals = [n for n, _ in nums]
            mx = max(nums, key=lambda t: t[0]); mn = min(nums, key=lambda t: t[0])
            lab = lambda ri: (data[ri][label_col] if label_col != ci and data[ri][label_col] else f"satır {first_row + ri}")
            lines.append(f"{h}: {len(vals)} değer; toplam {_fmt(sum(vals))}; ortalama {_fmt(sum(vals) / len(vals))}; "
                         f"en küçük {_fmt(mn[0])} ({lab(mn[1])}); en büyük {_fmt(mx[0])} ({lab(mx[1])}).")
        elif filled:
            from collections import Counter
            c = Counter(filled)
            if len(c) <= 25:
                top = ", ".join(f"{k} ({n})" for k, n in c.most_common(8))
                lines.append(f"{h}: {len(c)} farklı değer — {top}.")
            else:
                lines.append(f"{h}: {len(c)} farklı değer (ör. {', '.join(list(c)[:5])}).")
    return "\n".join(lines)


def _xlsx(data: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages: list[dict] = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else (v.strftime("%d.%m.%Y") if hasattr(v, "strftime") else
                                              (f"{v:g}" if isinstance(v, float) else str(v))) for v in r])
            if len(rows) > 20000:
                break
        pages += _table_pages(ws.title, rows, len(pages) + 1)
    title = (wb.properties.title or "").strip() or None if wb.properties else None
    return pages, title


def _csv(data: bytes):
    text = _decode(data)
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except Exception:  # noqa
        dialect = csv.excel
        if sample.count(";") > sample.count(","):
            dialect = type("d", (csv.excel,), {"delimiter": ";"})
    rows = list(csv.reader(io.StringIO(text), dialect))[:20000]
    return _table_pages("Tablo", rows, 1), None


def _pptx(data: bytes):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        texts, title = [], None
        try:
            if slide.shapes.title is not None and slide.shapes.title.text.strip():
                title = slide.shapes.title.text.strip()
        except Exception:  # noqa
            pass
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                t = sh.text_frame.text.strip()
                if t != title:
                    texts.append(t)
            if getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells))
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:  # noqa
            pass
        body = "\n".join(texts)
        if notes:
            body += f"\n[Konuşmacı notu] {notes}"
        full = _clean(f"{title}\n{body}" if title else body)
        pages.append({"page_number": i, "title": title or f"Slayt {i}", "text": full or f"Slayt {i}"})
    return pages, None


def html_to_sections(html: str) -> tuple[list[tuple[str | None, str]], str | None, dict]:
    """HTML'den ana metni (menu/reklam haric) basliklariyla cikarir."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "lxml")
    meta = {}
    title = None
    for sel, attr in (('meta[property="og:title"]', "content"), ('meta[name="title"]', "content")):
        el = soup.select_one(sel)
        if el and el.get(attr):
            title = el[attr].strip(); break
    if not title and soup.title and soup.title.string:
        title = soup.title.string.strip()
    for sel, key in (('meta[name="author"]', "author"), ('meta[property="article:published_time"]', "date"),
                     ('meta[name="description"]', "description"), ('meta[property="og:site_name"]', "site")):
        el = soup.select_one(sel)
        if el and el.get("content"):
            meta[key] = el["content"].strip()
    for t in soup(["script", "style", "noscript", "nav", "header", "footer", "aside", "form", "iframe",
                   "svg", "button", "figure"]):
        t.decompose()
    for t in soup.select('[class*="cookie"],[class*="banner"],[class*="advert"],[class*="share"],'
                         '[class*="related"],[class*="comment"],[id*="comment"],[class*="sidebar"],[role="navigation"]'):
        t.decompose()
    # en cok paragraf metni tasiyan kapsayici = ana icerik
    cands = soup.select("article, main, [role=main]") or []
    best, best_len = None, 0
    for c in cands + soup.find_all(["div", "section"]):
        L = sum(len(p.get_text(" ", strip=True)) for p in c.find_all("p", recursive=True))
        if L > best_len * 1.15 or (c.name in ("article", "main") and L >= best_len * 0.8 and L > 400):
            best, best_len = c, L
    root = best or soup.body or soup
    lines: list[tuple[int, str]] = []
    for el in root.find_all(["h1", "h2", "h3", "h4", "p", "li", "blockquote", "pre", "tr", "dt", "dd"]):
        txt = el.get_text(" ", strip=True)
        if not txt:
            continue
        if el.name in ("h1", "h2", "h3", "h4"):
            lines.append((int(el.name[1]), txt))
        elif el.name == "tr":
            lines.append((0, " | ".join(c.get_text(" ", strip=True) for c in el.find_all(["td", "th"]))))
        elif el.name == "li" and el.find(["p", "li"]):
            continue
        else:
            lines.append((0, txt))
    return _by_headings(lines), title, meta


def _html(data: bytes):
    secs, title, _ = html_to_sections(_decode(data))
    return _to_pages(secs), title


def _epub(data: bytes):
    z = zipfile.ZipFile(io.BytesIO(data))
    names = z.namelist()
    # okuma sirasi: content.opf spine
    order = []
    try:
        from bs4 import BeautifulSoup
        cont = BeautifulSoup(z.read("META-INF/container.xml"), "xml")
        opf_path = cont.find("rootfile")["full-path"]
        base = opf_path.rsplit("/", 1)[0] + "/" if "/" in opf_path else ""
        opf = BeautifulSoup(z.read(opf_path), "xml")
        items = {i["id"]: i["href"] for i in opf.find_all("item")}
        for ref in opf.find_all("itemref"):
            href = items.get(ref.get("idref"))
            if href:
                order.append(base + href.split("#")[0])
        t = opf.find("dc:title") or opf.find("title")
        title = t.get_text(strip=True) if t else None
    except Exception:  # noqa
        title = None
    if not order:
        order = sorted(n for n in names if n.lower().endswith((".xhtml", ".html", ".htm")))
    secs = []
    for n in order:
        if n not in names:
            continue
        s, t, _ = html_to_sections(_decode(z.read(n)))
        secs += s
    return _to_pages(secs), title


_READERS = {"docx": _docx, "md": _markdown, "txt": _txt, "rtf": _rtf, "xlsx": _xlsx, "csv": _csv,
            "pptx": _pptx, "epub": _epub, "html": _html}


def extract(kind: str, data: bytes) -> tuple[list[dict], str | None]:
    fn = _READERS.get(kind)
    if not fn:
        raise AppError("Bu dosya türü desteklenmiyor.")
    try:
        pages, title = fn(data)
    except AppError:
        raise
    except Exception as e:  # noqa
        raise AppError(f"Dosya okunamadı ({kind}). Bozuk ya da parola korumalı olabilir.") from e
    if not pages or sum(len(p["text"]) for p in pages) < 20:
        raise AppError("Dosyada okunabilir metin bulunamadı.")
    return pages, title


def text_pages(text: str) -> list[dict]:
    """Yapistirilan metin: markdown basliklari varsa ona gore, yoksa uzunluga gore bolumler."""
    if re.search(r"^#{1,3}\s", text, re.M):
        return _markdown(text.encode("utf-8"))[0]
    return _to_pages([(None, text)])
